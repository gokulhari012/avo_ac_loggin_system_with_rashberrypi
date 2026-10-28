"""Audit reports, status classification, offline detection and audit immutability."""

import io
from datetime import datetime, timedelta

import pytest
from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.exc import DatabaseError

from app.extensions import db
from app.models import AuditEvent, CommunicationEvent, Controller, ExportRecord, Node, NodeHeartbeat
from app.models.audit import AuditImmutableError
from app.services.audit import audit
from app.services.monitor import purge_heartbeats, run_monitor_cycle
from app.services.reports import ReportParams, build_downtime, build_excursions, build_missing_data
from app.services.status import controller_status
from tests.conftest import comm_failure, make_node, node_headers, reading, utcnow

BASE = datetime(2026, 3, 1, 12, 0, 0)


def ingest(client, readings):
    r = client.post("/api/v1/readings", json={"readings": readings}, headers=node_headers())
    assert r.status_code == 200 and not r.json["rejected"], r.json


def params(**kw):
    return ReportParams(start=BASE - timedelta(hours=1), end=BASE + timedelta(hours=1), **kw)


# ------------------------------------------------------------------ reports
def test_excursions(client, app):
    make_node()
    temps = [-18, -16, -14, -13, -16, -18, -26, -27, -19]
    ingest(client, [reading(at=BASE + timedelta(seconds=10 * i), temp=t) for i, t in enumerate(temps)])
    report = build_excursions(params(high_c=-15.0, low_c=-25.0))
    rows = report.sections[0].rows
    assert [(r[2], r[7], r[8]) for r in rows] == [("high", -13.0, 2), ("low", -27.0, 2)]
    assert report.sections[0].basis == "CALCULATED"


def test_excursion_requires_a_limit(app):
    with pytest.raises(ValueError):
        build_excursions(params())


def test_missing_data_report(client, app):
    make_node()
    ingest(client, [reading(at=BASE), reading(at=BASE + timedelta(seconds=10)),
                    comm_failure(at=BASE + timedelta(seconds=20)), comm_failure(at=BASE + timedelta(seconds=30)),
                    reading(at=BASE + timedelta(seconds=40)), reading(at=BASE + timedelta(minutes=20))])
    Controller.query.one().poll_interval_s = 10
    db.session.commit()
    report = build_missing_data(params())
    gaps = [r for r in report.sections[0].rows if r[5] == "no reading attempts recorded"]
    assert len(gaps) == 1 and gaps[0][4].startswith("19m")
    invalid = report.sections[1].rows
    assert len(invalid) == 1 and invalid[0][5] == 2 and invalid[0][6] == "timeout"


def test_downtime_inferred_from_heartbeat_gaps(app):
    node, _ = make_node()
    node.created_at = BASE - timedelta(days=1)
    for minutes in (0, 1, 2, 30, 31):
        db.session.add(NodeHeartbeat(node_id=node.id, received_at=BASE + timedelta(minutes=minutes)))
    db.session.commit()
    report = build_downtime(ReportParams(start=BASE, end=BASE + timedelta(minutes=32)))
    inferred = report.sections[0].rows
    assert report.sections[0].basis == "INFERRED"
    assert [r[4] for r in inferred] == ["heartbeat gap"] and inferred[0][3].startswith("28m")


@pytest.mark.parametrize("report_type", ["excursions", "faults", "downtime", "missing_data", "config_changes",
                                         "user_activity", "export_activity", "audit_trail"])
@pytest.mark.parametrize("fmt", ["xlsx", "pdf"])
def test_every_report_generates(as_user, users, client, report_type, fmt):
    make_node()
    ingest(client, [reading(at=utcnow() - timedelta(minutes=5), faults=["door_open"], fault_flags=16),
                    reading(at=utcnow() - timedelta(minutes=4))])
    c = as_user("admin")
    r = c.post("/reports/generate", data={"report_type": report_type, "format": fmt,
                                          "start": (utcnow() - timedelta(days=1)).strftime("%Y-%m-%dT%H:%M"),
                                          "end": (utcnow() + timedelta(hours=1)).strftime("%Y-%m-%dT%H:%M"),
                                          "high_c": "-15", "low_c": "-25"})
    assert r.status_code == 200, r.data[:300]
    if fmt == "pdf":
        assert r.data.startswith(b"%PDF")
    else:
        wb = load_workbook(io.BytesIO(r.data))
        assert wb.sheetnames[0] == "Report"
    record = ExportRecord.query.filter_by(report_type=report_type).one()
    assert record.outcome == "success" and record.sha256


def test_audit_reports_need_view_audit(as_user, users):
    r = as_user("operator").post("/reports/generate", data={"report_type": "audit_trail", "format": "pdf",
                                                            "start": "2026-01-01T00:00", "end": "2026-01-02T00:00"})
    assert r.status_code == 403


def test_temperature_history_report_redirects_to_logs(as_user, users):
    r = as_user("operator").post("/reports/generate", data={"report_type": "temperature_history",
                                                            "start": "2026-01-01T00:00", "end": "2026-01-02T00:00"})
    assert r.status_code == 302 and "/logs" in r.headers["Location"]


# ------------------------------------------------------------------ status
def make_ctrl(client, **reading_kwargs):
    make_node()
    ingest(client, [reading(**reading_kwargs)])
    node = Node.query.one()
    node.last_seen_at = utcnow()
    db.session.commit()
    return Controller.query.one(), node


def test_status_ok(client, app):
    ctrl, node = make_ctrl(client)
    assert controller_status(ctrl, node).code == "ok"


def test_status_stale(client, app):
    ctrl, node = make_ctrl(client, at=utcnow() - timedelta(minutes=10))
    assert controller_status(ctrl, node).code == "stale"


def test_status_node_offline(client, app):
    ctrl, node = make_ctrl(client)
    node.last_seen_at = utcnow() - timedelta(hours=1)
    assert controller_status(ctrl, node).code == "node_offline"


def test_status_comm_failure(client, app):
    make_node()
    ingest(client, [reading(at=utcnow() - timedelta(seconds=30)), comm_failure()])
    ctrl, node = Controller.query.one(), Node.query.one()
    node.last_seen_at = utcnow()
    assert controller_status(ctrl, node).code == "comm_failure"


def test_status_probe_fault(client, app):
    ctrl, node = make_ctrl(client, temp=None, quality="probe_fault", faults=["probe_fail_high"], fault_flags=2)
    assert controller_status(ctrl, node).code == "probe_fault"


def test_status_invalid(client, app):
    ctrl, node = make_ctrl(client, temp=None, quality="invalid", quality_flags=["out_of_range"])
    assert controller_status(ctrl, node).code == "invalid"


def test_status_fault(client, app):
    ctrl, node = make_ctrl(client, faults=["door_open"], fault_flags=16)
    status = controller_status(ctrl, node)
    assert status.code == "fault" and "door_open" in status.detail


def test_status_no_data(app):
    node, _ = make_node()
    node.last_seen_at = utcnow()
    assert controller_status(Controller.query.one(), node).code == "no_data"


def test_dashboard_marks_stale_reading_not_live(as_user, users, client):
    make_ctrl(client, at=utcnow() - timedelta(minutes=10))
    data = as_user("operator").get("/ui-api/dashboard").json
    card = data["controllers"][0]
    assert card["live"] is False and card["status"]["code"] == "stale"


def test_monitor_cycle_marks_offline_and_logs_events(client, app):
    node, _ = make_node()
    client.post("/api/v1/heartbeat", json={}, headers=node_headers())
    node = Node.query.one()
    assert node.online
    result = run_monitor_cycle(now=utcnow() + timedelta(minutes=10))
    assert result["went_offline"] == ["NODE-1"]
    assert CommunicationEvent.query.filter_by(event_type="node_offline").count() == 1
    assert AuditEvent.query.filter_by(action="node_offline").count() == 1
    client.post("/api/v1/heartbeat", json={}, headers=node_headers())
    assert AuditEvent.query.filter_by(action="node_online").count() == 2


def test_heartbeat_retention(app):
    node, _ = make_node()
    db.session.add_all([NodeHeartbeat(node_id=node.id, received_at=utcnow() - timedelta(days=200)),
                        NodeHeartbeat(node_id=node.id, received_at=utcnow())])
    db.session.commit()
    assert purge_heartbeats() == 1 and NodeHeartbeat.query.count() == 1


# ------------------------------------------------------------------ audit immutability
def test_audit_events_cannot_be_modified_via_orm(app):
    event = audit("test", "created", actor_type="system", commit=True)
    event.action = "tampered"
    with pytest.raises(AuditImmutableError):
        db.session.commit()
    db.session.rollback()
    db.session.delete(AuditEvent.query.first())
    with pytest.raises(AuditImmutableError):
        db.session.commit()
    db.session.rollback()


def test_audit_events_cannot_be_modified_via_sql(app):
    audit("test", "created", actor_type="system", commit=True)
    with pytest.raises(DatabaseError):
        db.session.execute(text("UPDATE audit_events SET action='x'"))
    db.session.rollback()
    with pytest.raises(DatabaseError):
        db.session.execute(text("DELETE FROM audit_events"))
    db.session.rollback()
    assert AuditEvent.query.count() == 1


def test_audit_redacts_secrets(app):
    event = audit("test", "x", actor_type="system", previous={"token": "abc", "password": "p"},
                  new={"token_hash": "h", "name": "ok"}, commit=True)
    assert "abc" not in event.previous_value and "[redacted]" in event.previous_value
    assert event.new == {"token_hash": "[redacted]", "name": "ok"}
