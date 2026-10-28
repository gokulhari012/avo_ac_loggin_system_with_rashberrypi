"""Audit and operational reports.

Every section states its basis:
  RECORDED   - rows stored as they happened (audit events, fault transitions,
               node-reported events, configuration versions, exports)
  CALCULATED - derived by the server from recorded readings (excursions, gaps)
  INFERRED   - estimated from the absence of data (node downtime from missing
               heartbeats); the true cause is unknown
No report claims compliance with any regulatory standard.
"""

from __future__ import annotations

import os
from dataclasses import dataclass, field
from datetime import datetime, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.platypus import Paragraph, Spacer

from ..extensions import db
from ..models import (AuditEvent, CommunicationEvent, Controller, ControllerConfigHistory, ExportRecord, FaultEvent,
                      Node, NodeConfigVersion, NodeHeartbeat, TemperatureReading)
from . import pdf_layout as P
from .exports import ExportRejected, fail_export, finish_export, start_export, temp_path
from .series import find_gaps, gap_threshold
from .settings import get_setting
from .timeutil import display_zone, duration_text, fmt_local, iso_utc, utcnow

RECORDED, CALCULATED, INFERRED = "RECORDED", "CALCULATED", "INFERRED"
MAX_ROWS = 50_000

REPORT_TYPES: dict[str, dict] = {
    "excursions": {"title": "High / low temperature excursions", "permission": "export_data"},
    "faults": {"title": "Controller faults and alarms", "permission": "export_data"},
    "downtime": {"title": "Node downtime and communication interruptions", "permission": "export_data"},
    "missing_data": {"title": "Missing data periods", "permission": "export_data"},
    "config_changes": {"title": "Configuration changes", "permission": "view_audit"},
    "user_activity": {"title": "User activity", "permission": "view_audit"},
    "export_activity": {"title": "Data export activity", "permission": "view_audit"},
    "audit_trail": {"title": "Audit trail", "permission": "view_audit"},
}


@dataclass
class Section:
    title: str
    basis: str
    columns: list[str]
    rows: list[list]
    note: str = ""
    widths: list[float] | None = None


@dataclass
class ReportData:
    report_type: str
    title: str
    params: list[tuple[str, str]]
    summary: list[tuple[str, str]] = field(default_factory=list)
    sections: list[Section] = field(default_factory=list)


@dataclass
class ReportParams:
    start: datetime
    end: datetime
    node_id: int | None = None
    controller_id: int | None = None
    high_c: float | None = None
    low_c: float | None = None
    category: str | None = None
    actor: str | None = None
    outcome: str | None = None

    def to_json(self) -> dict:
        return {k: (iso_utc(v) if isinstance(v, datetime) else v) for k, v in self.__dict__.items() if v is not None}

    def describe(self) -> list[tuple[str, str]]:
        out = [("Period", f"{fmt_local(self.start)} – {fmt_local(self.end)} ({display_zone().key})")]
        if self.node_id:
            node = db.session.get(Node, self.node_id)
            out.append(("Node", node.label if node else str(self.node_id)))
        if self.controller_id:
            ctrl = db.session.get(Controller, self.controller_id)
            out.append(("Controller", f"{ctrl.node.uid}/{ctrl.label}" if ctrl else str(self.controller_id)))
        for label, value in (("High limit", self.high_c), ("Low limit", self.low_c)):
            if value is not None:
                out.append((label, f"{value} °C"))
        for label, value in (("Category", self.category), ("Actor", self.actor), ("Outcome", self.outcome)):
            if value:
                out.append((label, value))
        return out


def _controllers(p: ReportParams) -> list[Controller]:
    q = Controller.query.join(Node)
    if p.controller_id:
        q = q.filter(Controller.id == p.controller_id)
    if p.node_id:
        q = q.filter(Controller.node_id == p.node_id)
    return q.order_by(Node.uid, Controller.uid).all()


def _nodes(p: ReportParams) -> list[Node]:
    q = Node.query
    if p.node_id:
        q = q.filter(Node.id == p.node_id)
    return q.order_by(Node.uid).all()


def _t(dt) -> str:
    return fmt_local(dt) if dt else ""


# ------------------------------------------------------------------ builders
def build_excursions(p: ReportParams) -> ReportData:
    if p.high_c is None and p.low_c is None:
        raise ValueError("enter a high limit, a low limit, or both")
    rows, totals = [], {"high": 0, "low": 0}
    for ctrl in _controllers(p):
        threshold = gap_threshold(ctrl)
        q = (db.session.query(TemperatureReading.acquired_at, TemperatureReading.temperature_c)
             .filter(TemperatureReading.controller_id == ctrl.id, TemperatureReading.acquired_at >= p.start,
                     TemperatureReading.acquired_at <= p.end, TemperatureReading.temperature_c.isnot(None))
             .order_by(TemperatureReading.acquired_at))
        current = None
        previous_t = None
        for t, value in q.yield_per(5000):
            kind = "high" if p.high_c is not None and value > p.high_c else (
                "low" if p.low_c is not None and value < p.low_c else None)
            if current and (kind != current["kind"]):
                current["end"] = t
                rows.append(_excursion_row(ctrl, current, p))
                current = None
            if kind and current is None:
                current = {"kind": kind, "start": t, "peak": value, "count": 0, "gap": False, "end": None}
            if current:
                if previous_t is not None and t - previous_t > threshold:
                    current["gap"] = True
                current["count"] += 1
                current["peak"] = max(current["peak"], value) if kind == "high" else min(current["peak"], value)
                current["last"] = t
            previous_t = t
        if current:
            rows.append(_excursion_row(ctrl, current, p))
    for r in rows:
        totals[r[2]] += 1
    return ReportData("excursions", REPORT_TYPES["excursions"]["title"], p.describe(),
                      [("Excursions above high limit", str(totals["high"])),
                       ("Excursions below low limit", str(totals["low"]))],
                      [Section("Excursions", CALCULATED,
                               ["Node", "Controller", "Type", "Limit °C", "Started", "Ended", "Duration",
                                "Peak °C", "Readings", "Notes"], rows[:MAX_ROWS],
                               "An excursion starts at the first valid reading beyond the limit and ends at the first "
                               "valid reading back within it. Readings without a valid temperature are ignored; "
                               "'data gap inside' means readings were missing during the excursion, so its true "
                               "extent is uncertain.")])


def _excursion_row(ctrl, e, p) -> list:
    limit = p.high_c if e["kind"] == "high" else p.low_c
    end = e["end"]
    notes = []
    if end is None:
        notes.append("still beyond limit at end of period")
    if e["gap"]:
        notes.append("data gap inside")
    duration = ((end or e["last"]) - e["start"]).total_seconds()
    return [ctrl.node.uid, ctrl.label, e["kind"], limit, _t(e["start"]), _t(end), duration_text(duration),
            round(e["peak"], 2), e["count"], "; ".join(notes)]


def build_faults(p: ReportParams) -> ReportData:
    ids = [c.id for c in _controllers(p)]
    events = (FaultEvent.query.filter(FaultEvent.controller_id.in_(ids), FaultEvent.started_at <= p.end,
                                      (FaultEvent.cleared_at.is_(None)) | (FaultEvent.cleared_at >= p.start))
              .order_by(FaultEvent.started_at).limit(MAX_ROWS).all()) if ids else []
    counts: dict[str, int] = {}
    rows = []
    for fe in events:
        counts[fe.code] = counts.get(fe.code, 0) + 1
        end = fe.cleared_at or utcnow()
        rows.append([fe.controller.node.uid, fe.controller.label, fe.kind, fe.code, fe.description,
                     _t(fe.started_at), _t(fe.cleared_at) or "still active", duration_text((end - fe.started_at)),
                     "" if fe.start_temperature_c is None else fe.start_temperature_c])
    return ReportData("faults", REPORT_TYPES["faults"]["title"], p.describe(),
                      [(f"{code}", str(n)) for code, n in sorted(counts.items())] or [("Fault events", "0")],
                      [Section("Fault and alarm periods", RECORDED,
                               ["Node", "Controller", "Kind", "Code", "Description", "Started", "Cleared",
                                "Duration", "Temp at start °C"], rows,
                               "Periods are bounded by transitions of the controller's fault status register (40012) "
                               "and alarm relay bit (40011 bit 1) in recorded readings. While the controller could "
                               "not be read the fault state is unknown and no transition is recorded.")])


def build_downtime(p: ReportParams) -> ReportData:
    threshold = timedelta(seconds=get_setting("node_offline_after_s"))
    inferred = []
    for node in _nodes(p):
        beats = [b for (b,) in db.session.query(NodeHeartbeat.received_at)
                 .filter(NodeHeartbeat.node_id == node.id, NodeHeartbeat.received_at >= p.start - threshold,
                         NodeHeartbeat.received_at <= p.end).order_by(NodeHeartbeat.received_at)]
        start_ref = max(p.start, node.created_at)
        if not beats:
            if start_ref < p.end:
                inferred.append([node.uid, _t(start_ref), _t(p.end), duration_text(p.end - start_ref),
                                 "no heartbeat in the whole period"])
            continue
        if beats[0] - start_ref > threshold:
            inferred.append([node.uid, _t(start_ref), _t(beats[0]), duration_text(beats[0] - start_ref),
                             "no heartbeat at start of period"])
        for gap in find_gaps(beats, threshold):
            inferred.append([node.uid, _t(gap.start), _t(gap.end), duration_text(gap.seconds), "heartbeat gap"])
        end_ref = min(p.end, utcnow())
        if end_ref - beats[-1] > threshold:
            inferred.append([node.uid, _t(beats[-1]), _t(end_ref), duration_text(end_ref - beats[-1]),
                             "no heartbeat until end of period"])
    node_ids = [n.id for n in _nodes(p)]
    types = ("node_offline", "node_online", "comm_lost", "comm_restored", "server_connection_restored", "data_gap",
             "api_auth_failure", "node_started")
    events = (CommunicationEvent.query.filter(CommunicationEvent.node_id.in_(node_ids),
                                              CommunicationEvent.occurred_at >= p.start,
                                              CommunicationEvent.occurred_at <= p.end,
                                              CommunicationEvent.event_type.in_(types))
              .order_by(CommunicationEvent.occurred_at).limit(MAX_ROWS).all()) if node_ids else []
    if p.controller_id:
        events = [e for e in events if e.controller_id in (None, p.controller_id)]
    recorded = [[e.node.uid, e.controller.label if e.controller else "", e.event_type, e.severity, _t(e.occurred_at),
                 e.origin, e.message, _event_detail(e)] for e in events]
    return ReportData("downtime", REPORT_TYPES["downtime"]["title"], p.describe(),
                      [("Inferred downtime periods", str(len(inferred))), ("Recorded communication events",
                                                                           str(len(recorded))),
                       ("Heartbeat retention (days)", str(get_setting("heartbeat_retention_days")))],
                      [Section("Node downtime inferred from missing heartbeats", INFERRED,
                               ["Node", "From", "To", "Duration", "Basis"], inferred,
                               f"A node is considered down when no heartbeat arrived for more than "
                               f"{int(threshold.total_seconds())} s. The cause (power, network, server) cannot be "
                               "determined from this alone. Data buffered on the node during such periods is "
                               "normally backfilled - check the missing-data report."),
                       Section("Communication events", RECORDED,
                               ["Node", "Controller", "Event", "Severity", "Occurred", "Origin", "Message",
                                "Details"], recorded,
                               "node_offline/node_online are recorded by the server; the other events are reported "
                               "by the node (comm_lost/comm_restored = Modbus link to the controller, "
                               "server_connection_restored = network/server outage seen by the node).")])


def _event_detail(e: CommunicationEvent) -> str:
    d = e.details
    keys = ("outage_started_at", "outage_ended_at", "failed_attempts", "backlog_readings", "from", "to",
            "records_lost", "reason", "comm_status", "lost_since", "failed_polls", "last_seen_at")
    return "; ".join(f"{k}={d[k]}" for k in keys if d.get(k) not in (None, ""))


def build_missing_data(p: ReportParams) -> ReportData:
    gap_rows, invalid_rows = [], []
    for ctrl in _controllers(p):
        threshold = gap_threshold(ctrl)
        q = (db.session.query(TemperatureReading.acquired_at, TemperatureReading.temperature_c,
                              TemperatureReading.comm_status, TemperatureReading.quality)
             .filter(TemperatureReading.controller_id == ctrl.id, TemperatureReading.acquired_at >= p.start,
                     TemperatureReading.acquired_at <= p.end).order_by(TemperatureReading.acquired_at))
        previous, invalid_start, invalid_reason, invalid_count, last_t = None, None, set(), 0, None
        first = None
        for t, value, comm, quality in q.yield_per(5000):
            first = first or t
            if previous is not None and t - previous > threshold:
                gap_rows.append([ctrl.node.uid, ctrl.label, _t(previous), _t(t), duration_text(t - previous),
                                 "no reading attempts recorded"])
            if value is None:
                if invalid_start is None:
                    invalid_start, invalid_reason, invalid_count = t, set(), 0
                invalid_reason.add(comm if comm != "ok" else quality)
                invalid_count += 1
            elif invalid_start is not None:
                invalid_rows.append([ctrl.node.uid, ctrl.label, _t(invalid_start), _t(t),
                                     duration_text(t - invalid_start), invalid_count, ", ".join(sorted(invalid_reason))])
                invalid_start = None
            previous = last_t = t
        if invalid_start is not None:
            invalid_rows.append([ctrl.node.uid, ctrl.label, _t(invalid_start), _t(last_t) + " (end of data)",
                                 duration_text(last_t - invalid_start), invalid_count, ", ".join(sorted(invalid_reason))])
        if first is None:
            gap_rows.append([ctrl.node.uid, ctrl.label, _t(p.start), _t(p.end), duration_text(p.end - p.start),
                             "no readings in the whole period"])
        else:
            if first - p.start > threshold:
                gap_rows.append([ctrl.node.uid, ctrl.label, _t(p.start), _t(first), duration_text(first - p.start),
                                 "no readings at start of period"])
            end_ref = min(p.end, utcnow())
            if end_ref - previous > threshold:
                gap_rows.append([ctrl.node.uid, ctrl.label, _t(previous), _t(end_ref), duration_text(end_ref - previous),
                                 "no readings until end of period"])
    return ReportData("missing_data", REPORT_TYPES["missing_data"]["title"], p.describe(),
                      [("Gaps without readings", str(len(gap_rows))),
                       ("Periods without a valid temperature", str(len(invalid_rows)))],
                      [Section("Periods without reading attempts", CALCULATED,
                               ["Node", "Controller", "From", "To", "Duration", "Basis"], gap_rows[:MAX_ROWS],
                               f"Gap = more than {get_setting('gap_factor')} x the controller's poll interval "
                               "(minimum 30 s) between consecutive recorded reading attempts."),
                       Section("Periods with readings but no valid temperature", CALCULATED,
                               ["Node", "Controller", "From", "To", "Duration", "Readings", "Reason"],
                               invalid_rows[:MAX_ROWS],
                               "The node polled the controller but no valid temperature was obtained "
                               "(communication failure, probe fault or invalid value).")])


def _audit_query(p: ReportParams, **filters):
    q = AuditEvent.query.filter(AuditEvent.occurred_at >= p.start, AuditEvent.occurred_at <= p.end)
    if filters.get("categories"):
        q = q.filter(AuditEvent.category.in_(filters["categories"]))
    if filters.get("actor_type"):
        q = q.filter(AuditEvent.actor_type == filters["actor_type"])
    if p.category:
        q = q.filter(AuditEvent.category == p.category)
    if p.actor:
        q = q.filter(AuditEvent.actor_name == p.actor)
    if p.outcome:
        q = q.filter(AuditEvent.outcome == p.outcome)
    if p.node_id:
        node = db.session.get(Node, p.node_id)
        if node:
            q = q.filter(AuditEvent.resource_id.like(f"{node.uid}%"))
    return q.order_by(AuditEvent.occurred_at).limit(MAX_ROWS)


def _audit_rows(events) -> list[list]:
    return [[_t(e.occurred_at), e.event_uid, f"{e.actor_type}:{e.actor_name}", e.category, e.action,
             f"{e.resource_type or ''}:{e.resource_id or ''}".strip(":"), e.outcome, e.reason or "",
             _short(e.previous_value), _short(e.new_value), e.source_ip or "", e.request_id or ""] for e in events]


AUDIT_COLUMNS = ["Time", "Event ID", "Actor", "Category", "Action", "Resource", "Outcome", "Reason", "Previous",
                 "New", "Source IP", "Request ID"]


def _short(text: str | None, limit: int = 400) -> str:
    if not text:
        return ""
    return text if len(text) <= limit else text[:limit] + "…"


def build_config_changes(p: ReportParams) -> ReportData:
    q = NodeConfigVersion.query.filter(NodeConfigVersion.created_at >= p.start, NodeConfigVersion.created_at <= p.end)
    if p.node_id:
        q = q.filter(NodeConfigVersion.node_id == p.node_id)
    versions = q.order_by(NodeConfigVersion.created_at).limit(MAX_ROWS).all()
    version_rows = [[v.node.uid, v.version, _t(v.created_at), v.created_by.username if v.created_by else "system",
                     v.reason, v.status, _t(v.delivered_at), _t(v.acked_at), _short(v.ack_errors)] for v in versions]
    hq = ControllerConfigHistory.query.filter(ControllerConfigHistory.recorded_at >= p.start,
                                              ControllerConfigHistory.recorded_at <= p.end)
    ids = [c.id for c in _controllers(p)]
    hq = hq.filter(ControllerConfigHistory.controller_id.in_(ids)) if ids else hq.filter(False)
    history = [[h.controller.node.uid, h.controller.label, _t(h.recorded_at), h.source, h.actor,
                _short(h.changed_keys), _short(h.previous_json), _short(h.new_json), h.reason]
               for h in hq.order_by(ControllerConfigHistory.recorded_at).limit(MAX_ROWS)]
    audits = _audit_rows(_audit_query(p, categories=("config", "controller", "settings", "node", "command")))
    return ReportData("config_changes", REPORT_TYPES["config_changes"]["title"], p.describe(),
                      [("Node configuration versions", str(len(version_rows))),
                       ("Controller configuration records", str(len(history))),
                       ("Related audit events", str(len(audits)))],
                      [Section("Node configuration versions and acknowledgement", RECORDED,
                               ["Node", "Version", "Created", "Created by", "Reason", "Status", "Delivered",
                                "Acknowledged", "Node errors"], version_rows),
                       Section("Controller configuration history", RECORDED,
                               ["Node", "Controller", "Recorded", "Source", "Actor", "Changed", "Previous", "New",
                                "Reason"], history,
                               "source=server_config: changed by a user on the server. source=device_snapshot: "
                               "parameter values read from the controller (40031-40068); a change there was made on "
                               "the controller itself. The keypad password (PA) is never recorded."),
                       Section("Audit events (configuration, settings, nodes, commands)", RECORDED, AUDIT_COLUMNS,
                               audits)])


def build_user_activity(p: ReportParams) -> ReportData:
    rows = _audit_rows(_audit_query(p, actor_type="user"))
    logins = sum(1 for r in rows if r[4] == "login")
    failures = sum(1 for r in rows if r[4] == "login_failed")
    return ReportData("user_activity", REPORT_TYPES["user_activity"]["title"], p.describe(),
                      [("Events", str(len(rows))), ("Logins", str(logins)), ("Failed logins", str(failures))],
                      [Section("User actions", RECORDED, AUDIT_COLUMNS, rows)])


def build_export_activity(p: ReportParams) -> ReportData:
    q = ExportRecord.query.filter(ExportRecord.created_at >= p.start, ExportRecord.created_at <= p.end)
    if p.actor:
        q = q.filter(ExportRecord.username == p.actor)
    rows = [[_t(e.created_at), e.export_uid, e.username, e.report_type, e.format, e.outcome,
             "" if e.record_count is None else e.record_count, _t(e.date_from), _t(e.date_to),
             _short(e.filters_json, 300), e.file_name or "", e.sha256 or "", e.error or ""]
            for e in q.order_by(ExportRecord.created_at).limit(MAX_ROWS)]
    return ReportData("export_activity", REPORT_TYPES["export_activity"]["title"], p.describe(),
                      [("Exports", str(len(rows)))],
                      [Section("Exports", RECORDED,
                               ["Time", "Export ID", "User", "Report", "Format", "Outcome", "Records", "From", "To",
                                "Filters", "File", "SHA-256", "Error"], rows,
                               "SHA-256 is the hash of the generated file and can be used to verify a copy.")])


def build_audit_trail(p: ReportParams) -> ReportData:
    rows = _audit_rows(_audit_query(p))
    return ReportData("audit_trail", REPORT_TYPES["audit_trail"]["title"], p.describe(),
                      [("Events", str(len(rows)))], [Section("Audit events", RECORDED, AUDIT_COLUMNS, rows)])


BUILDERS = {"excursions": build_excursions, "faults": build_faults, "downtime": build_downtime,
            "missing_data": build_missing_data, "config_changes": build_config_changes,
            "user_activity": build_user_activity, "export_activity": build_export_activity,
            "audit_trail": build_audit_trail}


# ------------------------------------------------------------------ rendering
def render_report(report_type: str, params: ReportParams, fmt: str) -> tuple[str, str, object]:
    if report_type not in BUILDERS or fmt not in ("xlsx", "pdf"):
        raise ValueError("unknown report type or format")
    record = start_export(report_type, fmt, params.to_json(), params.start, params.end)
    try:
        report = BUILDERS[report_type](params)
    except ValueError as exc:
        fail_export(record, "rejected_invalid", str(exc))
        raise ExportRejected("rejected_invalid", str(exc)) from exc
    count = sum(len(s.rows) for s in report.sections)
    if any(len(s.rows) >= MAX_ROWS for s in report.sections):
        message = (f"A section reached the {MAX_ROWS:,}-row report limit. Narrow the period or filters; "
                   "nothing was truncated.")
        fail_export(record, "rejected_too_large", message)
        raise ExportRejected("rejected_too_large", message)
    path = temp_path("." + fmt)
    stamp = utcnow().strftime("%Y%m%d-%H%M%S")
    file_name = f"{report_type.replace('_', '-')}-{stamp}-{record.export_uid[:8]}.{fmt}"
    try:
        (_render_xlsx if fmt == "xlsx" else _render_pdf)(report, path, record)
    except Exception as exc:
        os.remove(path)
        fail_export(record, "failed", f"{type(exc).__name__}: {exc}")
        raise
    finish_export(record, path, count, file_name)
    return path, file_name, record


def _meta(report: ReportData, record) -> list[tuple[str, str]]:
    zone = display_zone()
    return [("Report", report.title), ("Report ID", record.export_uid),
            ("Generated", f"{fmt_local(record.created_at)} ({zone.key}) / {iso_utc(record.created_at)}"),
            ("Generated by", record.username)] + report.params + report.summary


def _render_xlsx(report: ReportData, path: str, record) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    ws["A1"] = report.title
    ws["A1"].font = Font(bold=True, size=14)
    row = 3
    for key, value in _meta(report, record):
        ws.cell(row=row, column=1, value=key).font = Font(bold=True)
        ws.cell(row=row, column=2, value=value)
        row += 1
    row += 1
    ws.cell(row=row, column=1, value="Basis of each section").font = Font(bold=True)
    for section in report.sections:
        row += 1
        ws.cell(row=row, column=1, value=section.title)
        ws.cell(row=row, column=2, value=f"{section.basis}. {section.note}")
        ws.cell(row=row, column=2).alignment = Alignment(wrap_text=True)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 110
    for index, section in enumerate(report.sections, 1):
        sheet = wb.create_sheet(f"{index}. {section.title}"[:31].replace("/", "-"))
        for col, header in enumerate(section.columns, 1):
            cell = sheet.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F3A5F")
            sheet.column_dimensions[cell.column_letter].width = max(12, min(45, len(header) + 6))
        for values in section.rows:
            sheet.append(values)
        sheet.freeze_panes = "A2"
        if section.rows:
            sheet.auto_filter.ref = sheet.dimensions
        if not section.rows:
            sheet.cell(row=2, column=1, value="No records in this period.")
    wb.save(path)


def _render_pdf(report: ReportData, path: str, record) -> None:
    zone = display_zone()
    doc = P.document(path, report.title)
    story = [Paragraph(P.esc(report.title), P.STYLES["title"]), P.kv_table(_meta(report, record)), Spacer(1, 6),
             P.note("No compliance with any regulatory standard is claimed by this report.")]
    for section in report.sections:
        story.append(Paragraph(f"{P.esc(section.title)} <font size=8 color='#555555'>[{section.basis}]</font>",
                               P.STYLES["h2"]))
        if section.note:
            story.append(P.note(section.note))
        rows = section.rows
        story += P.data_tables(section.columns, [["" if v is None else v for v in r] for r in rows[:5000]],
                               chunk=200)
        if len(rows) > 5000:
            story.append(P.note(f"{len(rows)} rows in total; the PDF lists the first 5000. The Excel version of this "
                                "report contains all rows."))
    doc.build(story, canvasmaker=P.numbered_canvas_factory(
        f"{report.title} · ID {record.export_uid} · generated {fmt_local(record.created_at)} {zone.key}"))
